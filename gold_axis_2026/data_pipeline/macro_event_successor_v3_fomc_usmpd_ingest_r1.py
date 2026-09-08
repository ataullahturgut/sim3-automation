from __future__ import annotations

import hashlib
import io
import json
import math
import os
import subprocess
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

import numpy as np
import openpyxl
import psycopg
from psycopg.rows import dict_row
import requests

ENGINE = "MACRO_EVENT_SUCCESSOR_V3"
PIPELINE_VERSION = "MACRO_EVENT_SUCCESSOR_V3_FOMC_USMPD_HISTORICAL_R1_2026-09-08"
USMPD_URL = "https://www.frbsf.org/wp-content/uploads/USMPD.xlsx"
USMPD_PAGE = "https://www.frbsf.org/research-and-insights/data-and-indicators/us-monetary-policy-event-study-database/"
CALIBRATION_CUTOFF = datetime(2016, 1, 1)
ET = ZoneInfo("America/New_York")
UTC = timezone.utc
FUTNAMES = ("MP1", "MP2", "ED2", "ED3", "ED4")
RAW_SERIES = {
    "MP1": "MACRO_FOMC_USMPD_MP1",
    "MP2": "MACRO_FOMC_USMPD_MP2",
    "ED2": "MACRO_FOMC_USMPD_ED2",
    "ED3": "MACRO_FOMC_USMPD_ED3",
    "ED4": "MACRO_FOMC_USMPD_ED4",
}
TARGET_ID = "MACRO_FOMC_GSS_TARGET_FROZEN2015"
PATH_ID = "MACRO_FOMC_GSS_PATH_FROZEN2015"
ALL_IDS = tuple(RAW_SERIES.values()) + (TARGET_ID, PATH_ID)
DECISION_TABLES = ("monthly_forecast_contracts", "decision_signal_snapshots", "decision_runs", "decision_events")


@dataclass(frozen=True)
class Event:
    event_time_local: datetime
    event_time_utc: datetime
    date_key: str
    sep: int
    unscheduled: int
    values: dict[str, float]


def db_url() -> str:
    value = os.environ.get("NEON_DATABASE_URL", "").strip()
    if not value:
        raise RuntimeError("NEON_DATABASE_URL_NOT_SET")
    return value


def git_sha() -> str | None:
    try:
        return subprocess.run(["git", "rev-parse", "HEAD"], check=True, capture_output=True, text=True, timeout=5).stdout.strip()
    except Exception:
        return os.environ.get("GITHUB_SHA")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def stable_hash(payload: dict) -> str:
    return hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str).encode()).hexdigest()


def decision_counts(conn) -> dict[str, int]:
    out = {}
    with conn.cursor(row_factory=dict_row) as cur:
        for table in DECISION_TABLES:
            cur.execute(f"select count(*)::bigint n from {table}")
            out[table] = int(cur.fetchone()["n"])
    return out


def download_usmpd() -> tuple[bytes, str, datetime]:
    retrieved = datetime.now(UTC)
    r = requests.get(USMPD_URL, headers={"User-Agent": "Gold-Control-Macro-Event-USMPD/1.0"}, timeout=(10, 60))
    r.raise_for_status()
    ct = str(r.headers.get("Content-Type") or "")
    if "spreadsheet" not in ct and not r.content.startswith(b"PK"):
        raise RuntimeError(f"USMPD_UNEXPECTED_CONTENT_TYPE:{ct}")
    return r.content, sha256_bytes(r.content), retrieved


def parse_events(content: bytes) -> tuple[list[Event], str | None]:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if "Statements" not in wb.sheetnames or "README" not in wb.sheetnames:
        raise RuntimeError(f"USMPD_REQUIRED_SHEETS_MISSING:{wb.sheetnames}")
    ws = wb["Statements"]
    rows = ws.iter_rows(values_only=True)
    header = [str(x) if x is not None else "" for x in next(rows)]
    required = ["Date", "date_time", "SEP", "Unscheduled", *FUTNAMES]
    missing = [x for x in required if x not in header]
    if missing:
        raise RuntimeError(f"USMPD_REQUIRED_COLUMNS_MISSING:{missing}")
    pos = {name: header.index(name) for name in required}
    out: list[Event] = []
    for row in rows:
        dt = row[pos["date_time"]]
        if dt is None:
            continue
        if not isinstance(dt, datetime):
            try:
                dt = datetime.fromisoformat(str(dt))
            except Exception as exc:
                raise RuntimeError(f"USMPD_UNPARSEABLE_DATETIME:{dt}") from exc
        dt_local = dt.replace(tzinfo=ET) if dt.tzinfo is None else dt.astimezone(ET)
        vals = {}
        complete = True
        for name in FUTNAMES:
            v = row[pos[name]]
            if v is None:
                complete = False
                break
            try:
                fv = float(v)
            except Exception:
                complete = False
                break
            if not math.isfinite(fv):
                complete = False
                break
            vals[name] = fv
        if not complete:
            continue
        out.append(Event(
            event_time_local=dt_local,
            event_time_utc=dt_local.astimezone(UTC),
            date_key=dt_local.date().isoformat(),
            sep=int(row[pos["SEP"]] or 0),
            unscheduled=int(row[pos["Unscheduled"]] or 0),
            values=vals,
        ))
    out.sort(key=lambda e: e.event_time_utc)
    if len(out) < 200:
        raise RuntimeError(f"USMPD_COMPLETE_STATEMENT_COUNT_TOO_LOW:{len(out)}")

    readme = wb["README"]
    last_update = None
    for row in readme.iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, str) and cell.lower().startswith("last update:"):
                last_update = cell.split(":", 1)[1].strip()
                break
        if last_update:
            break
    return out, last_update


def sample_sd(x: np.ndarray, axis=0) -> np.ndarray:
    return np.std(x, axis=axis, ddof=1)


def ols_slope_with_intercept(x: np.ndarray, y: np.ndarray) -> np.ndarray:
    X = np.column_stack([np.ones(len(x)), x]) if x.ndim == 1 else np.column_stack([np.ones(len(x)), x])
    beta, *_ = np.linalg.lstsq(X, y, rcond=None)
    return beta


def fit_gss(dev: list[Event]) -> dict:
    X = np.array([[e.values[n] for n in FUTNAMES] for e in dev], dtype=float)
    means = X.mean(axis=0)
    sds = sample_sd(X, axis=0)
    if np.any(~np.isfinite(sds)) or np.any(sds <= 0):
        raise RuntimeError(f"GSS_DEVELOPMENT_ZERO_OR_INVALID_SD:{sds.tolist()}")
    Zs = (X - means) / sds
    cov = np.cov(Zs, rowvar=False, ddof=1)
    eigvals, eigvecs = np.linalg.eigh(cov)
    order = np.argsort(eigvals)[::-1]
    eigvals = eigvals[order]
    eigvecs = eigvecs[:, order]
    if eigvals[1] <= 0:
        raise RuntimeError("GSS_SECOND_EIGENVALUE_NONPOSITIVE")
    W = eigvecs[:, :2]
    F = Zs @ W @ np.diag(1.0 / np.sqrt(eigvals[:2]))

    mp1 = X[:, 0]
    g, *_ = np.linalg.lstsq(F, mp1, rcond=None)  # no intercept, as in gss.R
    if g[0] < 0:
        F = -F
        W = -W
        g = -g
    if abs(g[0]) < 1e-12:
        raise RuntimeError("GSS_ROTATION_G1_ZERO")
    alpha1 = 1.0 / math.sqrt(1.0 + (g[1] / g[0]) ** 2)
    alpha2 = alpha1 * g[1] / g[0]
    if abs(alpha2) < 1e-12:
        raise RuntimeError("GSS_ROTATION_ALPHA2_ZERO")
    beta1 = 1.0 / math.sqrt(1.0 + (alpha1 / alpha2) ** 2)
    beta2 = -alpha1 / alpha2 * beta1
    U = np.array([[alpha1, beta1], [alpha2, beta2]], dtype=float)
    rot = F @ U

    b1 = ols_slope_with_intercept(rot[:, 0], mp1)
    target_scale = float(b1[1])
    target = rot[:, 0] * target_scale

    ed4 = X[:, 4]
    b2 = ols_slope_with_intercept(np.column_stack([target, rot[:, 1]]), ed4)
    if abs(b2[1]) < 1e-12:
        raise RuntimeError("GSS_ED4_TARGET_COEF_ZERO")
    path_scale = float(b2[2] / b2[1])
    path = rot[:, 1] * path_scale

    norm_mp1 = ols_slope_with_intercept(np.column_stack([target, path]), mp1)
    norm_ed4 = ols_slope_with_intercept(np.column_stack([target, path]), ed4)
    checks = {
        "mp1_target_coef": float(norm_mp1[1]),
        "mp1_path_coef": float(norm_mp1[2]),
        "ed4_target_coef": float(norm_ed4[1]),
        "ed4_path_coef": float(norm_ed4[2]),
        "target_path_ed4_effect_ratio": float(norm_ed4[2] / norm_ed4[1]) if abs(norm_ed4[1]) > 1e-12 else None,
    }
    if abs(checks["mp1_target_coef"] - 1.0) > 1e-8 or abs(checks["mp1_path_coef"]) > 1e-8:
        raise RuntimeError(f"GSS_TARGET_NORMALIZATION_CHECK_FAILED:{checks}")
    if abs(checks["target_path_ed4_effect_ratio"] - 1.0) > 1e-8:
        raise RuntimeError(f"GSS_PATH_NORMALIZATION_CHECK_FAILED:{checks}")

    return {
        "means": means,
        "sds": sds,
        "W": W,
        "eigvals": eigvals[:2],
        "U": U,
        "target_scale": target_scale,
        "path_scale": path_scale,
        "checks": checks,
        "development_count": len(dev),
        "development_start": dev[0].date_key,
        "development_end": dev[-1].date_key,
    }


def apply_gss(events: list[Event], fit: dict) -> list[dict]:
    out = []
    for e in events:
        x = np.array([e.values[n] for n in FUTNAMES], dtype=float)
        z = (x - fit["means"]) / fit["sds"]
        f = z @ fit["W"] @ np.diag(1.0 / np.sqrt(fit["eigvals"]))
        rot = f @ fit["U"]
        target = float(rot[0] * fit["target_scale"])
        path = float(rot[1] * fit["path_scale"])
        if not (math.isfinite(target) and math.isfinite(path)):
            raise RuntimeError(f"GSS_NONFINITE_FACTOR:{e.date_key}")
        out.append({"event": e, "target": target, "path": path})
    return out


def registry_rows(workbook_hash: str, last_update: str | None) -> list[tuple]:
    common = {
        "authority_page": USMPD_PAGE,
        "download_url": USMPD_URL,
        "workbook_sha256": workbook_hash,
        "last_update_text": last_update,
        "event_window": "Statements 30-minute high-frequency window",
        "evidence_class": "HISTORICAL_REPLAY_RECONSTRUCTION",
        "calibration_cutoff_exclusive": CALIBRATION_CUTOFF.date().isoformat(),
    }
    rows = []
    for field, sid in RAW_SERIES.items():
        rows.append((sid, f"FOMC_USMPD_{field}", "Federal Reserve Bank of San Francisco USMPD", field,
                     "AUTHORITY_PUBLIC", "event", "percentage_points", "FOMC source component",
                     "CURRENT_MACRO_EVENT_SUCCESSOR_V3_HISTORICAL_INPUT", None,
                     json.dumps({**common, "field": field}, sort_keys=True)))
    for sid, role, method in [
        (TARGET_ID, "FOMC target surprise factor", "GSS_TARGET_FROZEN_PRE2016"),
        (PATH_ID, "FOMC path surprise factor", "GSS_PATH_FROZEN_PRE2016"),
    ]:
        rows.append((sid, sid, "Derived from Federal Reserve Bank of San Francisco USMPD", method,
                     "AUTHORITY_DERIVED", "event", "percentage_points", role,
                     "CURRENT_MACRO_EVENT_SUCCESSOR_V3_HISTORICAL_INPUT", None,
                     json.dumps({**common, "method": method, "future_outcome_use": False}, sort_keys=True)))
    return rows


def build_observations(events_with_factors: list[dict], run_id: str, retrieved: datetime, workbook_hash: str, last_update: str | None) -> list[dict]:
    out = []
    for row in events_with_factors:
        e: Event = row["event"]
        common = {
            "engine": ENGINE,
            "family": "FOMC",
            "event_date": e.date_key,
            "event_time_local": e.event_time_local.isoformat(),
            "event_timezone": "America/New_York",
            "sep": e.sep,
            "unscheduled": e.unscheduled,
            "source_url": USMPD_URL,
            "source_page": USMPD_PAGE,
            "workbook_sha256": workbook_hash,
            "workbook_last_update_text": last_update,
            "evidence_class": "HISTORICAL_REPLAY_RECONSTRUCTION",
            "historical_reconstruction_not_original_retrieval": True,
            "available_as_of_not_backdated": True,
            "calibration_cutoff_exclusive": CALIBRATION_CUTOFF.date().isoformat(),
            "production_authority": False,
            "direction_vote": False,
        }
        for field, sid in RAW_SERIES.items():
            value = e.values[field]
            meta = {**common, "source_field": field, "transform": "USMPD_STATEMENT_WINDOW_CHANGE"}
            out.append({
                "run_id": run_id, "series_id": sid, "observation_ts": e.event_time_utc, "value": value,
                "source": "Federal Reserve Bank of San Francisco USMPD", "source_symbol": field,
                "provider_as_of": retrieved, "available_as_of": retrieved, "first_seen_at": retrieved, "retrieved_at": retrieved,
                "frequency": "event", "unit": "percentage_points", "transform": "USMPD_STATEMENT_WINDOW_CHANGE",
                "quality_status": "APPROVED_AUTHORITY_HISTORICAL_RECONSTRUCTION", "lineage_id": f"USMPD_{workbook_hash[:12]}_{e.date_key}_{field}",
                "payload_hash": workbook_hash, "metadata": meta,
            })
        for sid, value, transform in [
            (TARGET_ID, row["target"], "GSS_TARGET_FROZEN_PRE2016"),
            (PATH_ID, row["path"], "GSS_PATH_FROZEN_PRE2016"),
        ]:
            meta = {**common, "transform": transform, "method_authority": "SF Fed gss.R / GSS 2005", "factor_fit_uses_post2015_events": False}
            out.append({
                "run_id": run_id, "series_id": sid, "observation_ts": e.event_time_utc, "value": value,
                "source": "Derived from Federal Reserve Bank of San Francisco USMPD", "source_symbol": transform,
                "provider_as_of": retrieved, "available_as_of": retrieved, "first_seen_at": retrieved, "retrieved_at": retrieved,
                "frequency": "event", "unit": "percentage_points", "transform": transform,
                "quality_status": "APPROVED_AUTHORITY_METHOD_HISTORICAL_RECONSTRUCTION", "lineage_id": f"USMPD_{workbook_hash[:12]}_{e.date_key}_{transform}",
                "payload_hash": stable_hash({"workbook_sha256": workbook_hash, "event": e.event_time_utc.isoformat(), "value": value, "transform": transform}),
                "metadata": meta,
            })
    return out


def main() -> int:
    content, workbook_hash, retrieved = download_usmpd()
    events, last_update = parse_events(content)
    dev = [e for e in events if e.event_time_local.replace(tzinfo=None) < CALIBRATION_CUTOFF]
    if len(dev) < 100:
        raise RuntimeError(f"GSS_DEVELOPMENT_SAMPLE_TOO_SMALL:{len(dev)}")
    fit = fit_gss(dev)
    transformed = apply_gss(events, fit)
    eval_count = sum(1 for x in transformed if x["event"].event_time_local.replace(tzinfo=None) >= CALIBRATION_CUTOFF)
    if eval_count < 50:
        raise RuntimeError(f"GSS_EVALUATION_SAMPLE_TOO_SMALL:{eval_count}")

    with psycopg.connect(db_url(), autocommit=False) as conn:
        before = decision_counts(conn)
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute("""select run_id,metadata from retrieval_runs
                           where pipeline_version=%s and status='SUCCESS'
                             and metadata->>'workbook_sha256'=%s
                           order by finished_at desc limit 1""", (PIPELINE_VERSION, workbook_hash))
            existing = cur.fetchone()
        if existing:
            result = {
                "status": "PASS_REUSED_EXISTING_USMPD_SOURCE_RUN",
                "engine_id": ENGINE,
                "source_run_id": str(existing["run_id"]),
                "workbook_sha256": workbook_hash,
                "complete_statement_events": len(events),
                "development_events": len(dev),
                "evaluation_events": eval_count,
                "gss_normalization_checks": fit["checks"],
                "production_authority": False,
            }
            print(json.dumps(result, indent=2, sort_keys=True)); return 0

        run_id = str(uuid.uuid4())
        obs = build_observations(transformed, run_id, retrieved, workbook_hash, last_update)
        metadata = {
            "engine_id": ENGINE,
            "family": "FOMC",
            "authority": "Federal Reserve Bank of San Francisco USMPD",
            "source_url": USMPD_URL,
            "workbook_sha256": workbook_hash,
            "workbook_last_update_text": last_update,
            "evidence_class": "HISTORICAL_REPLAY_RECONSTRUCTION",
            "complete_statement_events": len(events),
            "development_events": len(dev),
            "evaluation_events": eval_count,
            "calibration_cutoff_exclusive": CALIBRATION_CUTOFF.date().isoformat(),
            "gss_method": "SF Fed gss.R / Gürkaynak-Sack-Swanson target-path; fit frozen on pre-2016 only",
            "gss_normalization_checks": fit["checks"],
            "model_score_run": False,
            "production_authority": False,
            "forecast_or_decision_write": False,
            "market_shock_changed": False,
        }
        with conn.cursor() as cur:
            cur.execute("""insert into retrieval_runs
                (run_id,started_at,git_sha,pipeline_version,trigger_type,status,observations_read,observations_written,notes,metadata)
                values(%s,%s,%s,%s,%s,'RUNNING',%s,0,%s,%s::jsonb)""",
                (run_id, retrieved, git_sha(), PIPELINE_VERSION, "macro_event_v3_fomc_usmpd_historical",
                 len(obs), "SF Fed USMPD historical FOMC source/factor ingest; no production authority", json.dumps(metadata, sort_keys=True)))
            cur.executemany("""insert into source_registry
                (series_id,semantic_id,source_name,source_symbol,source_tier,frequency,unit,model_role,status,license_note,metadata)
                values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::jsonb)
                on conflict (series_id) do update set semantic_id=excluded.semantic_id,source_name=excluded.source_name,
                  source_symbol=excluded.source_symbol,source_tier=excluded.source_tier,frequency=excluded.frequency,
                  unit=excluded.unit,model_role=excluded.model_role,status=excluded.status,license_note=excluded.license_note,
                  metadata=excluded.metadata,updated_at=now()""", registry_rows(workbook_hash, last_update))
            cur.executemany("""insert into observations
                (run_id,series_id,observation_ts,value,source,source_symbol,provider_as_of,available_as_of,first_seen_at,retrieved_at,
                 frequency,unit,transform,quality_status,lineage_id,payload_hash,metadata)
                values (%(run_id)s,%(series_id)s,%(observation_ts)s,%(value)s,%(source)s,%(source_symbol)s,%(provider_as_of)s,%(available_as_of)s,
                        %(first_seen_at)s,%(retrieved_at)s,%(frequency)s,%(unit)s,%(transform)s,%(quality_status)s,%(lineage_id)s,%(payload_hash)s,%(metadata)s::jsonb)""",
                [{**o, "metadata": json.dumps(o["metadata"], sort_keys=True)} for o in obs])
            cur.execute("""update retrieval_runs set finished_at=now(),status='SUCCESS',observations_written=%s,metadata=%s::jsonb where run_id=%s""",
                        (len(obs), json.dumps(metadata, sort_keys=True), run_id))
        after = decision_counts(conn)
        if after != before:
            raise RuntimeError(f"DECISION_AUTHORITY_STORE_CHANGED:{before}->{after}")
        conn.commit()

    result = {
        "status": "PASS_USMPD_FOMC_HISTORICAL_INGEST",
        "engine_id": ENGINE,
        "source_run_id": run_id,
        "workbook_sha256": workbook_hash,
        "workbook_last_update_text": last_update,
        "complete_statement_events": len(events),
        "development_events": len(dev),
        "evaluation_events": eval_count,
        "series_count": len(ALL_IDS),
        "observations_written": len(obs),
        "gss_normalization_checks": fit["checks"],
        "evidence_class": "HISTORICAL_REPLAY_RECONSTRUCTION",
        "production_authority": False,
        "decision_counts_before": before,
        "decision_counts_after": after,
    }
    with open("macro_event_v3_fomc_usmpd_ingest_r1_result.json", "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, sort_keys=True)
    print(json.dumps(result, indent=2, sort_keys=True)); return 0


if __name__ == "__main__":
    raise SystemExit(main())

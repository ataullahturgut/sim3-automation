from __future__ import annotations
from ftplib import FTP
import io, json, re
from pathlib import Path
from openpyxl import load_workbook

HOST="ftp.cmegroup.com"
DIR="/daily_volume"
DATES=["20230103","20240102","20260930"]
OUT=Path("gold_h3_flow_ftp_probe_out"); OUT.mkdir(exist_ok=True)

def norm(v):
    if v is None: return ""
    return re.sub(r"\s+"," ",str(v).strip())

res={"host":HOST,"dir":DIR,"listing_sample":[],"files":[]}
ftp=FTP(HOST,timeout=30)
ftp.login()
ftp.cwd(DIR)
listing=ftp.nlst()
res["listing_count"]=len(listing)
res["listing_sample"]=listing[:50]
for d in DATES:
    fn=f"daily_volume_{d}.xlsx"
    item={"date":d,"filename":fn,"exists":fn in listing}
    if item["exists"]:
        bio=io.BytesIO()
        ftp.retrbinary(f"RETR {fn}",bio.write)
        raw=bio.getvalue()
        item["bytes"]=len(raw)
        try:
            wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
            hits=[]; oi_hits=[]
            for ws in wb.worksheets:
                for idx,row in enumerate(ws.iter_rows(values_only=True),1):
                    vals=[norm(x) for x in row]
                    low=" | ".join(vals).lower()
                    if "open interest" in low:
                        oi_hits.append({"sheet":ws.title,"row":idx,"values":vals})
                    if "gold" in low or re.search(r"(^|\W)gc(\W|$)",low):
                        hits.append({"sheet":ws.title,"row":idx,"values":vals})
            item["sheets"]=wb.sheetnames
            item["oi_hits"]=oi_hits[:50]
            item["gold_hits"]=hits[:100]
        except Exception as e:
            item["parse_error"]=repr(e)
    res["files"].append(item)
ftp.quit()

(OUT/"probe.json").write_text(json.dumps(res,indent=2,default=str))
lines=["# CME ANONYMOUS FTP DAILY_VOLUME — GC OI PROBE","",f"- host: {HOST}",f"- dir: {DIR}",f"- listing count: {res.get('listing_count')}",""]
for x in res["files"]:
    lines += [f"## {x['date']}",f"- exists: {x['exists']}"]
    if x["exists"]:
        lines += [f"- bytes: {x.get('bytes')}",f"- sheets: {x.get('sheets')}",f"- OI hits: {len(x.get('oi_hits',[]))}",f"- Gold/GC hits: {len(x.get('gold_hits',[]))}","", "### OI rows"]
        for h in x.get("oi_hits",[])[:20]:
            lines.append(f"- {h['sheet']}!{h['row']}: "+" | ".join(h["values"]))
        lines += ["","### Gold/GC rows"]
        for h in x.get("gold_hits",[])[:30]:
            lines.append(f"- {h['sheet']}!{h['row']}: "+" | ".join(h["values"]))
    lines.append("")
(OUT/"RESULT.md").write_text("\n".join(lines))
print((OUT/"RESULT.md").read_text())

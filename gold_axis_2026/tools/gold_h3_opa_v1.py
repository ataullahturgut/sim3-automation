from __future__ import annotations

from ftplib import FTP
from concurrent.futures import ThreadPoolExecutor, as_completed
import io, json, math, re, time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

ROOT=Path(__file__).resolve().parents[2]
AX=ROOT/"gold_axis_2026"
PANEL=AX/"GOLD_H3_CLEAN_PROSPECTIVE_V1_FROZEN_RIFT_PANEL.csv"
V5=AX/"GOLD_H3_CLEAN_HELIOS_V5_DCE_PREDICTIONS_2026-10-03.csv"
SAGE=AX/"GOLD_H3_SAGE_SELECTIVE_V1_PREDICTIONS_2026-10-04.csv"

OUT_SRC=AX/"GOLD_H3_OPA_V1_SOURCE_2026-10-04.csv"
OUT_COV=AX/"GOLD_H3_OPA_V1_SOURCE_COVERAGE_2026-10-04.csv"
OUT_PRED=AX/"GOLD_H3_OPA_V1_PREDICTIONS_2026-10-04.csv"
OUT_GRID=AX/"GOLD_H3_OPA_V1_THRESHOLD_GRID_2026-10-04.csv"
OUT_JSON=AX/"GOLD_H3_OPA_V1_SUMMARY_2026-10-04.json"
OUT_MD=AX/"GOLD_H3_OPA_V1_RESULT_2026-10-04.md"

HOST="ftp.cmegroup.com"
DIR="/daily_volume"
DATE_MIN=pd.Timestamp("2022-01-01")
DATE_MAX=pd.Timestamp("2026-09-30")
SEED=20261004
THRESH_GRID=[0.35,0.40,0.45,0.50,0.55,0.60]
FEATURES=[
    "oi_log_ratio","oi_asym","d_oi_asym_1","d_oi_log_ratio_1",
    "call_dlog_oi_1","put_dlog_oi_1",
    "vol_log_ratio","vol_asym","d_vol_asym_1","d_vol_log_ratio_1",
    "momentum_x_oi_asym","momentum_x_d_oi_asym",
]

def norm(x):
    if x is None: return ""
    return re.sub(r"\s+"," ",str(x).strip())

def safe_float(x):
    try:
        v=float(x)
        return v if np.isfinite(v) else np.nan
    except Exception:
        return np.nan

def parse_og_xlsx(raw:bytes, trade_date:pd.Timestamp, filename:str):
    wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    ws=next((s for s in wb.worksheets if "by Product" in s.title),None)
    if ws is None:
        return {"trade_date":trade_date,"filename":filename,"parse_status":"NO_PRODUCT_SHEET"}

    header=None; header_row=None
    for i,row in enumerate(ws.iter_rows(values_only=True),start=1):
        vals=[norm(x) for x in row]
        low=[v.lower() for v in vals]
        if "commodity indicator" in low and any("open interest" in v for v in low):
            header=vals; header_row=i; break
    if not header:
        return {"trade_date":trade_date,"filename":filename,"parse_status":"NO_HEADER"}

    def idx(keys):
        for j,h in enumerate(header):
            hh=h.lower()
            if any(k in hh for k in keys): return j
        raise KeyError(keys)
    try:
        i_ex=idx(["exchange name"])
        i_ci=idx(["commodity indicator"])
        i_pd=idx(["product description"])
        i_fo=idx(["future/option indicator"])
        i_tv=idx(["total volume"])
        i_oi=idx(["open interest"])
    except Exception:
        return {"trade_date":trade_date,"filename":filename,"parse_status":"HEADER_MAP_FAIL","header":" | ".join(header)}

    got={}
    for row in ws.iter_rows(min_row=header_row+1,values_only=True):
        vals=list(row)
        def get(j): return vals[j] if j < len(vals) else None
        ci=norm(get(i_ci)).upper()
        pdsc=norm(get(i_pd)).upper()
        foi=norm(get(i_fo)).upper()
        exch=norm(get(i_ex)).upper()
        if ci=="OG" and foi=="O" and ("COMEX" in exch) and pdsc in ("GOLD CALL","GOLD PUT"):
            side="call" if pdsc=="GOLD CALL" else "put"
            got[side]={"vol":safe_float(get(i_tv)),"oi":safe_float(get(i_oi))}
    if "call" not in got or "put" not in got:
        return {"trade_date":trade_date,"filename":filename,"parse_status":"OG_CALL_PUT_NOT_FOUND"}
    vals=[got["call"]["vol"],got["call"]["oi"],got["put"]["vol"],got["put"]["oi"]]
    if not all(np.isfinite(v) and v>0 for v in vals):
        return {
            "trade_date":trade_date,"filename":filename,"parse_status":"MISSING_VALUE",
            "og_call_volume":got["call"]["vol"],"og_call_oi":got["call"]["oi"],
            "og_put_volume":got["put"]["vol"],"og_put_oi":got["put"]["oi"],
        }
    return {
        "trade_date":trade_date,"filename":filename,"parse_status":"PASS",
        "og_call_volume":got["call"]["vol"],"og_call_oi":got["call"]["oi"],
        "og_put_volume":got["put"]["vol"],"og_put_oi":got["put"]["oi"],
    }

def connect():
    ftp=FTP(HOST,timeout=45)
    ftp.login(); ftp.cwd(DIR)
    return ftp

def worker(chunk):
    ftp=connect(); rows=[]
    try:
        for d,fn in chunk:
            raw=None; last=None
            for attempt in range(3):
                try:
                    bio=io.BytesIO()
                    ftp.retrbinary(f"RETR {fn}",bio.write)
                    raw=bio.getvalue(); break
                except Exception as e:
                    last=repr(e)
                    try: ftp.quit()
                    except Exception: pass
                    time.sleep(.5+attempt)
                    ftp=connect()
            if raw is None:
                rows.append({"trade_date":d,"filename":fn,"parse_status":"FTP_FAIL","error":last}); continue
            try: rows.append(parse_og_xlsx(raw,d,fn))
            except Exception as e:
                rows.append({"trade_date":d,"filename":fn,"parse_status":"PARSE_EXCEPTION","error":repr(e)})
    finally:
        try: ftp.quit()
        except Exception: pass
    return rows

def acquire():
    ftp=connect(); listing=ftp.nlst()
    try: ftp.quit()
    except Exception: pass
    pat=re.compile(r"daily_volume_(\d{8})\.xlsx$",re.I)
    targets=[]
    for fn in listing:
        m=pat.search(fn)
        if not m: continue
        d=pd.to_datetime(m.group(1),format="%Y%m%d",errors="coerce")
        if pd.notna(d) and DATE_MIN<=d<=DATE_MAX: targets.append((d,fn))
    targets=sorted(set(targets))
    workers=10; chunks=[targets[i::workers] for i in range(workers)]
    rows=[]
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs=[ex.submit(worker,c) for c in chunks if c]
        for i,f in enumerate(as_completed(futs),1):
            rows.extend(f.result())
            print(f"worker {i}/{len(futs)} complete rows={len(rows)}")
    return pd.DataFrame(rows).sort_values("trade_date").reset_index(drop=True),len(listing)

def coverage(raw):
    x=raw.copy(); x["year"]=pd.to_datetime(x.trade_date).dt.year
    out=[]
    for y,g in x.groupby("year"):
        p=g[g.parse_status=="PASS"]; m=g[g.parse_status!="PASS"]
        out.append({
            "year":int(y),"listed_files":len(g),"pass_rows":len(p),"coverage":len(p)/max(len(g),1),
            "min_pass":str(p.trade_date.min().date()) if len(p) else None,
            "max_pass":str(p.trade_date.max().date()) if len(p) else None,
            "missing_rows":len(m),
            "missing_statuses":";".join(f"{k}:{v}" for k,v in m.parse_status.value_counts().to_dict().items())
        })
    return pd.DataFrame(out)

def source_features(raw):
    x=raw[raw.parse_status=="PASS"].copy()
    x=x.drop_duplicates("trade_date",keep="last").sort_values("trade_date").reset_index(drop=True)
    for c in ["og_call_volume","og_call_oi","og_put_volume","og_put_oi"]: x[c]=x[c].astype(float)
    eps=1e-12
    x["oi_log_ratio"]=np.log((x.og_call_oi+eps)/(x.og_put_oi+eps))
    x["oi_asym"]=(x.og_call_oi-x.og_put_oi)/(x.og_call_oi+x.og_put_oi+eps)
    x["d_oi_asym_1"]=x.oi_asym.diff()
    x["d_oi_log_ratio_1"]=x.oi_log_ratio.diff()
    x["call_dlog_oi_1"]=np.log(x.og_call_oi).diff()
    x["put_dlog_oi_1"]=np.log(x.og_put_oi).diff()
    x["vol_log_ratio"]=np.log((x.og_call_volume+eps)/(x.og_put_volume+eps))
    x["vol_asym"]=(x.og_call_volume-x.og_put_volume)/(x.og_call_volume+x.og_put_volume+eps)
    x["d_vol_asym_1"]=x.vol_asym.diff()
    x["d_vol_log_ratio_1"]=x.vol_log_ratio.diff()
    return x

def as_bool(s):
    if s.dtype==bool: return s
    return s.astype(str).str.lower().isin(["true","1","yes"])

def build_panel(src):
    p=pd.read_csv(PANEL)
    v=pd.read_csv(V5,usecols=["feature_cutoff_date","p_helios_v5_dce"])
    try:
        sg=pd.read_csv(SAGE,usecols=["feature_cutoff_date","ocs_candidate"])
        sg["ocs_candidate"]=as_bool(sg.ocs_candidate)
    except Exception:
        sg=pd.DataFrame(columns=["feature_cutoff_date","ocs_candidate"])
    for z in [p,v,sg]:
        z["feature_cutoff_date"]=pd.to_datetime(z.feature_cutoff_date)
    p["forecast_issue_date"]=pd.to_datetime(p.forecast_issue_date)
    p["target_end_date_h3"]=pd.to_datetime(p.target_end_date_h3)
    p=p.merge(v,on="feature_cutoff_date",how="inner",validate="one_to_one")
    p=p.merge(sg,on="feature_cutoff_date",how="left",validate="one_to_one")
    p["ocs_candidate"]=p.ocs_candidate.fillna(False).astype(bool)
    p["v5_pred"]=(p.p_helios_v5_dce>=.5).astype(int)
    p["eligible_v5_continuation"]=p.v5_pred.eq(p.momentum_up.astype(int))
    p["rescue_target"]=p.v5_pred.ne(p.y_up.astype(int)).astype(int)

    s=src.dropna(subset=[
        "oi_log_ratio","oi_asym","d_oi_asym_1","d_oi_log_ratio_1",
        "call_dlog_oi_1","put_dlog_oi_1","vol_log_ratio","vol_asym",
        "d_vol_asym_1","d_vol_log_ratio_1"
    ]).sort_values("trade_date").copy()
    rows=[]
    for r in p.itertuples():
        z=s[s.trade_date<r.feature_cutoff_date]
        if z.empty: continue
        q=z.iloc[-1]; stale=(r.feature_cutoff_date-q.trade_date).days
        if stale>7: continue
        d=r._asdict()
        for c in [
            "oi_log_ratio","oi_asym","d_oi_asym_1","d_oi_log_ratio_1",
            "call_dlog_oi_1","put_dlog_oi_1","vol_log_ratio","vol_asym",
            "d_vol_asym_1","d_vol_log_ratio_1"
        ]: d[c]=float(q[c])
        sign=1.0 if int(r.momentum_up)==1 else -1.0
        d["momentum_x_oi_asym"]=sign*float(q.oi_asym)
        d["momentum_x_d_oi_asym"]=sign*float(q.d_oi_asym_1)
        d["opa_trade_date"]=q.trade_date
        d["opa_stale_days"]=int(stale)
        rows.append(d)
    a=pd.DataFrame(rows)
    if not a.empty:
        a["month_key"]=a.forecast_issue_date.dt.to_period("M").astype(str)
    return a

def model():
    return Pipeline([
        ("scale",StandardScaler()),
        ("model",LogisticRegression(C=1.0,solver="lbfgs",max_iter=3000,class_weight="balanced",random_state=SEED))
    ])

def walk(panel):
    test=panel[(panel.forecast_issue_date>=pd.Timestamp("2023-01-01")) & panel.eligible_v5_continuation].copy()
    out=[]
    for mo in sorted(test.month_key.unique()):
        te=test[test.month_key==mo].copy()
        cutoff=te.feature_cutoff_date.min(); first_issue=te.forecast_issue_date.min()
        tr=panel[
            panel.eligible_v5_continuation
            & (panel.target_end_date_h3<=cutoff)
            & (panel.forecast_issue_date<first_issue)
        ].dropna(subset=FEATURES+["rescue_target"]).copy()
        if len(tr)<80 or tr.rescue_target.nunique()<2: continue
        m=model(); m.fit(tr[FEATURES].to_numpy(float),tr.rescue_target.astype(int).to_numpy())
        pp=m.predict_proba(te[FEATURES].to_numpy(float))[:,1]
        for r,pv in zip(te.itertuples(),pp):
            out.append({
                "feature_cutoff_date":r.feature_cutoff_date,"forecast_issue_date":r.forecast_issue_date,
                "target_end_date_h3":r.target_end_date_h3,"year":int(r.year),"month":str(r.month),
                "y_up":int(r.y_up),"target_r3":float(r.target_r3),"momentum_up":int(r.momentum_up),
                "p_v5":float(r.p_helios_v5_dce),"v5_pred":int(r.v5_pred),
                "rescue_target":int(r.rescue_target),"ocs_candidate":bool(r.ocs_candidate),
                "p_opa_rescue":float(pv),"opa_trade_date":r.opa_trade_date,
                "opa_stale_days":int(r.opa_stale_days),"train_n":int(len(tr)),
            })
    return pd.DataFrame(out).sort_values("forecast_issue_date").reset_index(drop=True)

def f2(p,r):
    return 0.0 if p<=0 or r<=0 else 5*p*r/(4*p+r)

def cand_stats(z,th):
    c=z.p_opa_rescue>=th; y=z.rescue_target.astype(bool)
    tp=int((c&y).sum()); fp=int((c&~y).sum()); fn=int((~c&y).sum())
    precision=tp/max(tp+fp,1); recall=tp/max(tp+fn,1); rate=float(c.mean()) if len(c) else 0.0
    return {
        "threshold":th,"eligible_n":len(z),"candidate_n":int(c.sum()),"true_rescue_n":int(y.sum()),
        "tp":tp,"fp":fp,"fn":fn,"precision":precision,"recall":recall,"candidate_rate":rate,
        "f2":f2(precision,recall),"eligible":bool(precision>=.45 and rate<=.35)
    }

def prob_metrics(y,p):
    y=np.asarray(y,int); p=np.clip(np.asarray(p,float),1e-6,1-1e-6)
    return {
        "accuracy":float(np.mean((p>=.5)==y)),
        "brier":float(np.mean((p-y)**2)),
        "logloss":float(np.mean(-(y*np.log(p)+(1-y)*np.log(1-p))))
    }

def period_eval(pred,th,year):
    q=pred[pred.year==year].copy()
    q["opa_candidate"]=q.p_opa_rescue>=th
    q["incremental_action"]=q.opa_candidate & (~q.ocs_candidate)
    q["sage_pred"]=np.where(q.ocs_candidate,1-q.v5_pred,q.v5_pred).astype(int)
    q["assisted_pred"]=np.where(q.incremental_action,1-q.v5_pred,q.sage_pred).astype(int)
    q["rescue"]=q.incremental_action & q.sage_pred.ne(q.y_up) & q.assisted_pred.eq(q.y_up)
    q["broken"]=q.incremental_action & q.sage_pred.eq(q.y_up) & q.assisted_pred.ne(q.y_up)
    nact=int(q.incremental_action.sum()); resc=int(q.rescue.sum()); brok=int(q.broken.sum())
    sage_p=np.where(q.ocs_candidate,1-q.p_v5,q.p_v5)
    assisted_p=np.where(q.incremental_action,1-q.p_v5,sage_p)
    sm=prob_metrics(q.y_up,sage_p); am=prob_metrics(q.y_up,assisted_p)
    half={}
    for h,mask in [("H1",q.forecast_issue_date.dt.month<=6),("H2",q.forecast_issue_date.dt.month>=7)]:
        z=q[mask]; rr=int(z.rescue.sum()); bb=int(z.broken.sum())
        half[h]={"n":len(z),"actions":int(z.incremental_action.sum()),"rescue":rr,"broken":bb,"net":rr-bb}
    return {
        "year":year,"covered_n":len(q),"opa_candidates":int(q.opa_candidate.sum()),
        "ocs_overlap":int((q.opa_candidate&q.ocs_candidate).sum()),"incremental_actions":nact,
        "rescue":resc,"broken":brok,"net":resc-brok,"action_precision":resc/max(nact,1),
        "sage_correct":int((q.sage_pred==q.y_up).sum()),"assisted_correct":int((q.assisted_pred==q.y_up).sum()),
        "sage_accuracy":float((q.sage_pred==q.y_up).mean()),"assisted_accuracy":float((q.assisted_pred==q.y_up).mean()),
        "sage_brier":sm["brier"],"assisted_brier":am["brier"],
        "sage_logloss":sm["logloss"],"assisted_logloss":am["logloss"],"half_year":half
    }

def main():
    raw,listing_count=acquire()
    raw.to_csv(OUT_SRC,index=False)
    cov=coverage(raw); cov.to_csv(OUT_COV,index=False)
    feat=source_features(raw)
    panel=build_panel(feat)
    if panel.empty: raise RuntimeError("NO_ALIGNED_PANEL")
    pred=walk(panel)
    if pred.empty: raise RuntimeError("NO_PREDICTIONS")

    dev=pred[pred.year.isin([2023,2024])].copy()
    grid=pd.DataFrame([cand_stats(dev,t) for t in THRESH_GRID])
    grid.to_csv(OUT_GRID,index=False)
    elig=grid[grid.eligible].copy()

    status="OPA_DEV_FAIL"; selected=None; confirm=None; hold=None
    if not elig.empty:
        elig=elig.sort_values(["f2","recall","precision","candidate_rate","threshold"],ascending=[False,False,False,True,False])
        selected=float(elig.iloc[0].threshold)
        ev25=period_eval(pred,selected,2025)
        confirm=bool(
            ev25["net"]>0 and ev25["action_precision"]>=.50 and ev25["rescue"]>=2
            and ev25["assisted_accuracy"]>=ev25["sage_accuracy"]
        )
        status="OPA_2025_CONFIRM_PASS" if confirm else "OPA_2025_CONFIRM_FAIL"
        if confirm:
            hold=period_eval(pred,selected,2026)
            status="OPA_2026_HOLDOUT_OPENED"

    pred.to_csv(OUT_PRED,index=False)
    summ={
        "schema":"OPA_H3_V1","status":status,"listing_count":listing_count,
        "source_rows":len(raw),"source_pass_rows":int((raw.parse_status=="PASS").sum()),
        "coverage":cov.to_dict("records"),"aligned_rows":len(panel),"prediction_rows":len(pred),
        "selected_threshold":selected,"dev_grid":grid.to_dict("records"),
        "confirmation_2025":period_eval(pred,selected,2025) if selected is not None else None,
        "confirmation_pass":confirm,"holdout_2026":hold
    }
    OUT_JSON.write_text(json.dumps(summ,indent=2,default=str)+"\n")

    lines=[
        "# OPA-H3 V1 — GOLD OPTIONS POSITIONING ASYMMETRY RESULT","",
        f"**Status:** **{status}**  ",
        "**Source:** official CME anonymous FTP daily_volume workbooks; OG GOLD CALL / GOLD PUT aggregate volume + preliminary OI.  ",
        "**PIT:** strict prior-trade-date only; same-day source values forbidden.","",
        "## Source coverage","",
        "| Year | Listed | PASS | Coverage | First | Last | Missing |",
        "|---:|---:|---:|---:|---|---|---:|"
    ]
    for r in cov.itertuples():
        lines.append(f"| {r.year} | {r.listed_files} | {r.pass_rows} | {100*r.coverage:.2f}% | {r.min_pass} | {r.max_pass} | {r.missing_rows} |")
    lines += ["","## DEV 2023-2024","",
              "| Threshold | Candidates | Precision | Recall | Rate | F2 | Eligible |",
              "|---:|---:|---:|---:|---:|---:|---|"]
    for r in grid.itertuples():
        lines.append(f"| {r.threshold:.2f} | {r.candidate_n} | {100*r.precision:.2f}% | {100*r.recall:.2f}% | {100*r.candidate_rate:.2f}% | {r.f2:.4f} | {r.eligible} |")
    if selected is not None:
        e=period_eval(pred,selected,2025)
        lines += ["",f"## Selected threshold: {selected:.2f}","","## 2025 untouched confirmation","",
                  f"- covered origins: **{e['covered_n']}**",
                  f"- OPA candidates / OCS overlap / incremental actions: **{e['opa_candidates']} / {e['ocs_overlap']} / {e['incremental_actions']}**",
                  f"- rescue / broken / net: **{e['rescue']} / {e['broken']} / {e['net']:+d}**",
                  f"- action precision: **{100*e['action_precision']:.2f}%**",
                  f"- SAGE V2 -> SAGE+OPA accuracy: **{100*e['sage_accuracy']:.2f}% -> {100*e['assisted_accuracy']:.2f}%**",
                  f"- confirmation: **{'PASS' if confirm else 'FAIL'}**"]
        if hold is not None:
            h=hold
            lines += ["","## 2026 final holdout","",
                      f"- covered origins: **{h['covered_n']}**",
                      f"- OPA candidates / OCS overlap / incremental actions: **{h['opa_candidates']} / {h['ocs_overlap']} / {h['incremental_actions']}**",
                      f"- rescue / broken / net: **{h['rescue']} / {h['broken']} / {h['net']:+d}**",
                      f"- action precision: **{100*h['action_precision']:.2f}%**",
                      f"- SAGE V2 -> SAGE+OPA accuracy: **{100*h['sage_accuracy']:.2f}% -> {100*h['assisted_accuracy']:.2f}%**",
                      f"- SAGE V2 -> assisted Brier: **{h['sage_brier']:.4f} -> {h['assisted_brier']:.4f}**",
                      f"- SAGE V2 -> assisted log-loss: **{h['sage_logloss']:.4f} -> {h['assisted_logloss']:.4f}**",
                      f"- H1 net: **{h['half_year']['H1']['net']:+d}**, H2 net: **{h['half_year']['H2']['net']:+d}**"]
    lines += ["","## Governance","",
              "2023-2024 selected the threshold. 2025 was the untouched confirmation. 2026 was evaluated only if the frozen 2025 gate passed. No 2026 outcome changed source alignment, features, model or threshold.",
              "OPA is options-positioning asymmetry, not implied-volatility skew and not CME CVOL."]
    OUT_MD.write_text("\n".join(lines)+"\n")
    print(OUT_MD.read_text())

if __name__=="__main__":
    main()

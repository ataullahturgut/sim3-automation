"""Extract only explicit late-2026 values from verified original data workbooks.

No hidden fallbacks, interpolation, source substitution, or as-of assumption.
The source files must be validated with original git blob/file gates before
running this extraction. Output is a *research-only* label/version manifest.
"""
from __future__ import annotations
import argparse,json,math,re
from datetime import datetime,date
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

RELEASES={
 "2026-09":{"release_utc":"2026-09-01T14:27:55Z",
            "commit":"21449bec606bc5a5b1bf58d08ce526dc35d66c94",
            "blob":"253056ac5d54881c1f3e5767087ef682e258c7a5"},
 "2026-10":{"release_utc":"2026-10-01T14:19:21Z",
            "commit":"8b903073ea4ce9caf9b2f238e792685ca17270f4",
            "blob":"a46ff84a7e9fe83f5d1b73186d161aeba57b4da0"}
}

def parse_period(x):
    if isinstance(x,(datetime,date)):return x.strftime("%Y-%m")
    s=str(x)
    if re.match(r"^20[12][0-9][Mm][01][0-9]$",s):
        return s[:4]+"-"+s[-2:]
    if re.match(r"^20[12][0-9]-[01][0-9]",s):
        return s[:7]
    return None

def gpr_extract(path,release):
    df=pd.read_excel(path,engine="xlrd",sheet_name="Sheet1")
    if not set(("month","GPR","GPRT","GPRA")).issubset(df.columns):
        raise RuntimeError("OFFICIAL_GPR_COLUMNS_DIFFER")
    out={}
    for _,r in df.iterrows():
        m=parse_period(r["month"])
        if m not in ("2026-07","2026-08","2026-09"):continue
        values={}
        for name in ("GPR","GPRT","GPRA"):
            v=float(r[name])
            if not (math.isfinite(v) and v>=0):raise RuntimeError(f"GPR_{name}_BAD_{m}")
            values[name]=v
        if m in out:raise RuntimeError(f"GPR_DUPLICATE_MONTH_{m}")
        out[m]=values
    if not out:raise RuntimeError("GPR_RECENT_MONTHS_MISSING")
    return {
        "vintage":release,**RELEASES[release],
        "original_git_path":f"gpr_archive_files/data_gpr_export_{release.replace('-','')}.xls",
        "series_values":out,
        "first_capture_by_our_system_at_release_verified":False,
        "current_archive_data_later_than_prediction_origin_must_be_excluded":True
    }

def wb_extract(path):
    w=load_workbook(path,read_only=True,data_only=True)
    if "Monthly Prices" not in w.sheetnames:raise RuntimeError("WB_MONTHLY_PRICES_MISSING")
    sheet=w["Monthly Prices"]
    banner=str(sheet.cell(1,1).value)
    if "World Bank" not in banner:raise RuntimeError("WB_SOURCE_BANNER_MISMATCH")
    header={c:sheet.cell(5,c).value for c in range(1,sheet.max_column+1)}
    goldcol=[c for c,v in header.items() if isinstance(v,str) and v.strip().lower()=="gold"]
    if len(goldcol)!=1:raise RuntimeError(f"WB_GOLD_COLUMN_AMBIGUOUS {goldcol}")
    col=goldcol[0]
    units=str(sheet.cell(6,col).value)
    if "oz" not in units.lower():raise RuntimeError(f"WB_GOLD_PRICE_UNIT_NOT_OZ {units}")
    update_note=str(sheet.cell(4,1).value)
    months={}
    for rr in sheet.iter_rows(min_row=7,values_only=True):
        if not rr:continue
        ym=parse_period(rr[0])
        if ym not in ("2026-07","2026-08","2026-09"):continue
        value=float(rr[col-1])
        if not (math.isfinite(value) and 0<value<100000):
            raise RuntimeError(f"WB_INVALID_GOLD_{ym}")
        if ym in months:raise RuntimeError(f"WB_DUPLICATE_GOLD_{ym}")
        months[ym]=value
    if "2026-08" not in months or "2026-09" not in months:
        raise RuntimeError(f"WB_LATEST_TWO_MONTHS_MISSING {months}")
    if "October 02, 2026" not in update_note:
        raise RuntimeError(f"WB_MONTHLY_PUBLISHED_DATE_NOT_CONFIRMED {update_note}")
    return {
       "source":"World Bank Commodity Price Data (Pink Sheet)",
       "sheet":"Monthly Prices","column":col,"unit":units,
       "published_date_note":update_note,
       "official_release_not_before":"2026-10-02",
       "gold_monthly_values_usd_oz":months,
       "not_available_to_forecast_same_month_at_prior_month_end":True,
       "original_live_capture_at_release_verified":False
    }

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--gpr-sep",required=True);p.add_argument("--gpr-oct",required=True)
    p.add_argument("--wb",required=True);p.add_argument("--report",required=True)
    a=p.parse_args()
    gpr={"2026-09":gpr_extract(a.gpr_sep,"2026-09"),
         "2026-10":gpr_extract(a.gpr_oct,"2026-10")}
    if "2026-08" not in gpr["2026-09"]["series_values"]:
        raise RuntimeError("GPR_AUG_MISSING_FROM_SEP_RELEASE")
    if "2026-09" not in gpr["2026-10"]["series_values"]:
        raise RuntimeError("GPR_SEP_MISSING_FROM_OCT_RELEASE")
    wb=wb_extract(a.wb)
    out={"schema":"GOLD_MONTHLY_OFFICIAL_GPR_WB_EXTRACT_20261023_V1",
         "status":"ORIGINAL_MONTHLY_VALUES_PARSED_AND_VERSIONED_RESEARCH_ONLY",
         "gpr":gpr,"world_bank":wb,
         "neon_written_by_parser":False,
         "can_auto_replace_prior_train_or_forecast":False}
    Path(a.report).write_text(json.dumps(out,indent=2,sort_keys=True)+chr(10))
    print("OFFICIAL_GPR_WB_MONTHLY_VALUES_EXTRACTION=PASS")
    print(json.dumps(out,sort_keys=True))
if __name__=="__main__":
    main()

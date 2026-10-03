from __future__ import annotations
import io, json, re, requests
from pathlib import Path
from openpyxl import load_workbook

DATES = ["20230103","20231229","20240102","20240328","20261002"]
BASE = "https://www.cmegroup.com/ftp/pub/pub/pub/daily_volume/daily_volume_{date}.xlsx"
OUT = Path("gold_h3_flow_daily_volume_probe_out"); OUT.mkdir(exist_ok=True)

def norm(v):
    if v is None: return ""
    return re.sub(r"\s+"," ",str(v).strip())

def fetch(date):
    url = BASE.format(date=date)
    r = requests.get(url, timeout=60, headers={"User-Agent":"Mozilla/5.0 research"})
    item={"date":date,"url":url,"status":r.status_code,"bytes":len(r.content),"content_type":r.headers.get("content-type")}
    if r.status_code != 200:
        item["head"]=r.text[:500]
        return item
    try:
        wb=load_workbook(io.BytesIO(r.content),data_only=True,read_only=True)
    except Exception as e:
        item["parse_error"]=repr(e)
        item["head_bytes"]=r.content[:100].hex()
        return item
    sheets=[]
    hits=[]
    oi_hits=[]
    for ws in wb.worksheets:
        rows=[]
        for idx,row in enumerate(ws.iter_rows(values_only=True),1):
            vals=[norm(x) for x in row]
            if any(vals):
                rows.append((idx,vals))
                low=" | ".join(vals).lower()
                if "gold" in low or re.search(r"(^|\W)gc(\W|$)",low):
                    hits.append({"sheet":ws.title,"row":idx,"values":vals})
                if "open interest" in low:
                    oi_hits.append({"sheet":ws.title,"row":idx,"values":vals})
        sheets.append({"sheet":ws.title,"rows":ws.max_row,"cols":ws.max_column,"sample":rows[:20]})
    item["sheets"]=sheets
    item["gold_hits"]=hits[:100]
    item["open_interest_text_hits"]=oi_hits[:100]
    return item

res=[fetch(d) for d in DATES]
(OUT/"probe.json").write_text(json.dumps(res,indent=2,default=str))
lines=["# CME PUBLIC DAILY_VOLUME XLSX — GC VOLUME/OI PROBE",""]
for x in res:
    lines += [f"## {x['date']}",f"- HTTP: {x['status']}",f"- bytes: {x['bytes']}",f"- content-type: {x.get('content_type')}"]
    if x["status"]==200 and "parse_error" not in x:
        lines += [f"- sheets: {', '.join(s['sheet'] for s in x['sheets'])}",f"- OI text hits: {len(x.get('open_interest_text_hits',[]))}",f"- Gold/GC hits: {len(x.get('gold_hits',[]))}",""]
        lines.append("### OI/header hits")
        for h in x.get("open_interest_text_hits",[])[:20]:
            lines.append(f"- {h['sheet']}!{h['row']}: " + " | ".join(h["values"]))
        lines.append("")
        lines.append("### Gold/GC hits")
        for h in x.get("gold_hits",[])[:40]:
            lines.append(f"- {h['sheet']}!{h['row']}: " + " | ".join(h["values"]))
    else:
        lines.append("- error/head: "+str(x.get("parse_error") or x.get("head")))
    lines.append("")
(OUT/"RESULT.md").write_text("\n".join(lines))
print((OUT/"RESULT.md").read_text())

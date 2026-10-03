from __future__ import annotations
import json, hashlib, io, re
from pathlib import Path
import requests
from openpyxl import load_workbook

DATES=["20220103","20230103","20260129"]
BASE="https://www.cmegroup.com/ftp/daily_volume/daily_volume_{date}.xlsx"
OUT=Path("gold_h3_flow_probe_out"); OUT.mkdir(exist_ok=True)

def norm(x):
    return re.sub(r"\s+"," ",str(x or "").strip())

def inspect(date):
    url=BASE.format(date=date)
    r=requests.get(url,timeout=60,headers={"User-Agent":"Mozilla/5.0 FLOW-H3 research probe"})
    r.raise_for_status()
    raw=r.content
    sha=hashlib.sha256(raw).hexdigest()
    wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    sheets=[]
    hits=[]
    for ws in wb.worksheets:
        rows=[]
        for i,row in enumerate(ws.iter_rows(values_only=True),start=1):
            vals=[norm(v) for v in row]
            if any(vals):
                rows.append((i,vals))
        sample=rows[:25]
        sheets.append({"sheet":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,"sample":sample})
        for i,vals in rows:
            joined=" | ".join(vals).lower()
            if "gold" in joined or "open interest" in joined or "volume" in joined:
                hits.append({"sheet":ws.title,"row":i,"values":vals})
    return {"date":date,"url":url,"status":r.status_code,"bytes":len(raw),"sha256":sha,"sheets":sheets,"hits":hits[:250]}

res=[inspect(d) for d in DATES]
(OUT/"probe.json").write_text(json.dumps(res,indent=2,default=str)+"\n")
lines=["# CME DAILY VOLUME XLSX — FLOW-H3 STRUCTURE PROBE",""]
for x in res:
    lines += [f"## {x['date']}",f"- bytes: {x['bytes']}",f"- sha256: {x['sha256']}",f"- sheets: {', '.join(s['sheet'] for s in x['sheets'])}","", "### Gold / volume / OI hits",""]
    for h in x["hits"][:80]:
        vals=" | ".join(v for v in h["values"] if v)
        lines.append(f"- {h['sheet']}!{h['row']}: {vals}")
    lines.append("")
(OUT/"RESULT.md").write_text("\n".join(lines)+"\n")
print((OUT/"RESULT.md").read_text())

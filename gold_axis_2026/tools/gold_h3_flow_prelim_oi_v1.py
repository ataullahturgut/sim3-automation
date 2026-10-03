from __future__ import annotations
from ftplib import FTP, error_temp
import io, json, math, re, time
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.metrics import log_loss

ROOT=Path(__file__).resolve().parents[2]
AX=ROOT/"gold_axis_2026"
PANEL=AX/"GOLD_H3_CLEAN_PROSPECTIVE_V1_FROZEN_RIFT_PANEL.csv"
V5=AX/"GOLD_H3_CLEAN_HELIOS_V5_DCE_PREDICTIONS_2026-10-03.csv"

OUT_SRC=AX/"GOLD_H3_FLOW_PRELIM_OI_V1_SOURCE_2026-10-03.csv"
OUT_COV=AX/"GOLD_H3_FLOW_PRELIM_OI_V1_SOURCE_COVERAGE_2026-10-03.csv"
OUT_PRED=AX/"GOLD_H3_FLOW_PRELIM_OI_V1_PREDICTIONS_2026-10-03.csv"
OUT_GRID=AX/"GOLD_H3_FLOW_PRELIM_OI_V1_THRESHOLD_GRID_2026-10-03.csv"
OUT_JSON=AX/"GOLD_H3_FLOW_PRELIM_OI_V1_SUMMARY_2026-10-03.json"
OUT_MD=AX/"GOLD_H3_FLOW_PRELIM_OI_V1_RESULT_2026-10-03.md"

HOST="ftp.cmegroup.com"
DIR="/daily_volume"
DATE_MIN=pd.Timestamp("2022-01-01")
DATE_MAX=pd.Timestamp("2026-09-30")
SEED=20261003
THRESH_GRID=[0.35,0.40,0.45,0.50,0.55]
FEATURES=[
    "dlog_volume_1","dlog_oi_1","volume_z20","oi_z20",
    "volume_oi_ratio","d_volume_oi_ratio_1","oi_accel_5","volume_accel_5",
    "momentum_x_dlog_oi","momentum_x_volume_z20"
]

def norm(x):
    if x is None: return ""
    return re.sub(r"\s+"," ",str(x).strip())

def safe_float(x):
    try:
        v=float(x)
        return v if np.isfinite(v) else np.nan
    except Exception:
        return np.nan

def parse_gc_xlsx(raw:bytes, trade_date:pd.Timestamp, filename:str):
    wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    ws=None
    for s in wb.worksheets:
        if "by Product" in s.title:
            ws=s; break
    if ws is None:
        return {"trade_date":trade_date,"filename":filename,"parse_status":"NO_PRODUCT_SHEET"}

    header=None; header_row=None
    for i,row in enumerate(ws.iter_rows(values_only=True),start=1):
        vals=[norm(x) for x in row]
        low=[v.lower() for v in vals]
        if "commodity indicator" in low and any("open interest" in v for v in low):
            header=vals; header_row=i; break
    if not header:
        return {"trade_date":trade_date,"filename":filename,"parse_status":"NO_HEADER"}

    def col_idx(keys, required=True):
        for j,h in enumerate(header):
            hh=h.lower()
            if any(k in hh for k in keys):
                return j
        if required: raise KeyError(keys)
        return None

    try:
        i_ex=col_idx(["exchange name"])
        i_ci=col_idx(["commodity indicator"])
        i_pd=col_idx(["product description"])
        i_fo=col_idx(["future/option indicator"])
        i_tv=col_idx(["total volume"])
        i_oi=col_idx(["open interest"])
    except Exception:
        return {"trade_date":trade_date,"filename":filename,"parse_status":"HEADER_MAP_FAIL","header":" | ".join(header)}

    for i,row in enumerate(ws.iter_rows(min_row=header_row+1,values_only=True),start=header_row+1):
        vals=list(row)
        def get(idx):
            return vals[idx] if idx is not None and idx < len(vals) else None
        ci=norm(get(i_ci)).upper()
        pdsc=norm(get(i_pd)).upper()
        foi=norm(get(i_fo)).upper()
        exch=norm(get(i_ex)).upper()
        if ci=="GC" and foi=="F" and ("GOLD FUTURES" in pdsc) and ("COMEX" in exch):
            vol=safe_float(get(i_tv))
            oi=safe_float(get(i_oi))
            status="PASS" if (np.isfinite(vol) and vol>0 and np.isfinite(oi) and oi>0) else "MISSING_VALUE"
            return {
                "trade_date":trade_date,"filename":filename,"parse_status":status,
                "exchange":exch,"commodity_indicator":ci,"product_description":pdsc,
                "gc_total_volume":vol,"gc_prelim_open_interest":oi
            }
    return {"trade_date":trade_date,"filename":filename,"parse_status":"GC_ROW_NOT_FOUND"}

def connect():
    ftp=FTP(HOST,timeout=45)
    ftp.login()
    ftp.cwd(DIR)
    return ftp

def acquire():
    ftp=connect()
    listing=ftp.nlst()
    pat=re.compile(r"daily_volume_(\d{8})\.xlsx$",re.I)
    targets=[]
    for fn in listing:
        m=pat.search(fn)
        if not m: continue
        d=pd.to_datetime(m.group(1),format="%Y%m%d",errors="coerce")
        if pd.isna(d): continue
        if DATE_MIN <= d <= DATE_MAX:
            targets.append((d,fn))
    targets=sorted(set(targets))

    rows=[]
    for k,(d,fn) in enumerate(targets,1):
        raw=None
        last=None
        for attempt in range(3):
            try:
                bio=io.BytesIO()
                ftp.retrbinary(f"RETR {fn}",bio.write)
                raw=bio.getvalue()
                break
            except Exception as e:
                last=repr(e)
                try: ftp.quit()
                except Exception: pass
                time.sleep(1+attempt)
                ftp=connect()
        if raw is None:
            rows.append({"trade_date":d,"filename":fn,"parse_status":"FTP_FAIL","error":last})
            continue
        try:
            rows.append(parse_gc_xlsx(raw,d,fn))
        except Exception as e:
            rows.append({"trade_date":d,"filename":fn,"parse_status":"PARSE_EXCEPTION","error":repr(e)})
        if k % 100 == 0:
            print(f"parsed {k}/{len(targets)}")
    try: ftp.quit()
    except Exception: pass
    return pd.DataFrame(rows).sort_values("trade_date").reset_index(drop=True), len(listing)

def source_features(src):
    x=src[src.parse_status=="PASS"].copy()
    x=x.drop_duplicates("trade_date",keep="last").sort_values("trade_date").reset_index(drop=True)
    x["logv"]=np.log(x.gc_total_volume.astype(float))
    x["logoi"]=np.log(x.gc_prelim_open_interest.astype(float))
    x["dlog_volume_1"]=x.logv.diff()
    x["dlog_oi_1"]=x.logoi.diff()

    vm=x.logv.shift(1).rolling(20,min_periods=10).mean()
    vs=x.logv.shift(1).rolling(20,min_periods=10).std(ddof=0)
    om=x.logoi.shift(1).rolling(20,min_periods=10).mean()
    os=x.logoi.shift(1).rolling(20,min_periods=10).std(ddof=0)
    x["volume_z20"]=(x.logv-vm)/vs.replace(0,np.nan)
    x["oi_z20"]=(x.logoi-om)/os.replace(0,np.nan)
    x["volume_oi_ratio"]=x.gc_total_volume/x.gc_prelim_open_interest
    x["d_volume_oi_ratio_1"]=x.volume_oi_ratio.pct_change()
    x["oi_accel_5"]=x.dlog_oi_1-x.dlog_oi_1.shift(1).rolling(5,min_periods=3).mean()
    x["volume_accel_5"]=x.dlog_volume_1-x.dlog_volume_1.shift(1).rolling(5,min_periods=3).mean()
    return x

def coverage(raw):
    raw=raw.copy()
    raw["year"]=pd.to_datetime(raw.trade_date).dt.year
    out=[]
    for y,g in raw.groupby("year"):
        p=g[g.parse_status=="PASS"]
        miss=g[g.parse_status!="PASS"]
        out.append({
            "year":int(y),"listed_files":len(g),"pass_rows":len(p),
            "coverage":len(p)/max(len(g),1),
            "min_pass":str(p.trade_date.min().date()) if len(p) else None,
            "max_pass":str(p.trade_date.max().date()) if len(p) else None,
            "missing_rows":len(miss),
            "missing_statuses":";".join(f"{k}:{v}" for k,v in miss.parse_status.value_counts().to_dict().items())
        })
    return pd.DataFrame(out)

def align(panel,src):
    p=panel.copy()
    p["feature_cutoff_date"]=pd.to_datetime(p.feature_cutoff_date)
    p["forecast_issue_date"]=pd.to_datetime(p.forecast_issue_date)
    p["target_end_date_h3"]=pd.to_datetime(p.target_end_date_h3)
    s=src.dropna(subset=FEATURES).sort_values("trade_date").copy()
    rows=[]
    for r in p.itertuples():
        z=s[s.trade_date < r.feature_cutoff_date]  # preregistered strict prior-date
        if z.empty: continue
        q=z.iloc[-1]
        stale=(r.feature_cutoff_date-q.trade_date).days
        if stale>7:  # calendar staleness guard; holidays/weekends tolerated, long gaps fail closed
            continue
        sign=1.0 if int(r.momentum_up)==1 else -1.0
        d=r._asdict()
        d["flow_trade_date"]=q.trade_date
        d["flow_stale_days"]=int(stale)
        d["flow_volume"]=float(q.gc_total_volume)
        d["flow_oi"]=float(q.gc_prelim_open_interest)
        for c in ["dlog_volume_1","dlog_oi_1","volume_z20","oi_z20","volume_oi_ratio",
                  "d_volume_oi_ratio_1","oi_accel_5","volume_accel_5"]:
            d[c]=float(q[c])
        d["momentum_x_dlog_oi"]=sign*float(q.dlog_oi_1)
        d["momentum_x_volume_z20"]=sign*float(q.volume_z20)
        rows.append(d)
    a=pd.DataFrame(rows)
    if not a.empty:
        a["month_key"]=a.forecast_issue_date.dt.to_period("M").astype(str)
    return a

def model():
    return Pipeline([
        ("scale",StandardScaler()),
        ("model",LogisticRegression(C=1.0,solver="lbfgs",max_iter=3000,class_weight="balanced",random_state=SEED))
    ])

def walk(panel):
    test=panel[panel.forecast_issue_date>=pd.Timestamp("2023-01-01")].copy()
    out=[]
    for mo in sorted(test.month_key.unique()):
        te=test[test.month_key==mo].copy()
        cutoff=te.feature_cutoff_date.min()
        first_issue=te.forecast_issue_date.min()
        tr=panel[(panel.target_end_date_h3<=cutoff)&(panel.forecast_issue_date<first_issue)].copy()
        tr=tr.dropna(subset=FEATURES+["reversal_target"])
        if len(tr)<80 or tr.reversal_target.nunique()<2: continue
        m=model()
        m.fit(tr[FEATURES].to_numpy(float),tr.reversal_target.astype(int).to_numpy())
        pr=m.predict_proba(te[FEATURES].to_numpy(float))[:,1]
        for r,pv in zip(te.itertuples(),pr):
            out.append({
                "feature_cutoff_date":r.feature_cutoff_date,
                "forecast_issue_date":r.forecast_issue_date,
                "target_end_date_h3":r.target_end_date_h3,
                "year":int(r.year),"month":str(r.month),"y_up":int(r.y_up),
                "target_r3":float(r.target_r3),"p_aurora":float(r.p_aurora),
                "momentum_up":int(r.momentum_up),"reversal_target":int(r.reversal_target),
                "aurora_pred":int(r.aurora_pred),
                "aurora_follows_momentum":bool(r.aurora_follows_momentum),
                "flow_trade_date":r.flow_trade_date,
                "flow_stale_days":int(r.flow_stale_days),
                "flow_volume":float(r.flow_volume),
                "flow_oi":float(r.flow_oi),
                "p_flow_reversal":float(pv),"train_n":int(len(tr))
            })
    return pd.DataFrame(out).sort_values("forecast_issue_date").reset_index(drop=True)

def f2(p,r):
    return 0.0 if p<=0 or r<=0 else 5*p*r/(4*p+r)

def cand_metrics(z,th):
    e=z[z.aurora_follows_momentum.astype(bool)].copy()
    c=e.p_flow_reversal>=th
    y=e.reversal_target.astype(bool)
    tp=int((c&y).sum()); fp=int((c&~y).sum()); fn=int((~c&y).sum())
    precision=tp/max(tp+fp,1); recall=tp/max(tp+fn,1); rate=float(c.mean()) if len(c) else 0.0
    return {
        "threshold":th,"eligible_n":len(e),"candidate_n":int(c.sum()),"true_reversal_n":int(y.sum()),
        "tp":tp,"fp":fp,"fn":fn,"precision":precision,"recall":recall,"candidate_rate":rate,
        "f2":f2(precision,recall),"eligible":bool(precision>=.45 and rate<=.35)
    }

def parse_bool(s):
    if s.dtype==bool: return s
    return s.astype(str).str.lower().isin(["true","1","yes"])

def merge_opal(pred):
    v=pd.read_csv(V5)
    v["feature_cutoff_date"]=pd.to_datetime(v.feature_cutoff_date)
    keep=["feature_cutoff_date","p_helios_v5_dce","opal_override_check"]
    z=pred.merge(v[keep],on="feature_cutoff_date",how="left",validate="one_to_one")
    z["opal_candidate"]=parse_bool(z.opal_override_check)
    z["v5_pred"]=(z.p_helios_v5_dce>=.5).astype(int)
    return z

def period_stats(z,th,year):
    q=z[(z.year==year)&z.aurora_follows_momentum.astype(bool)].copy()
    q["flow_candidate"]=q.p_flow_reversal>=th
    y=q.reversal_target.astype(bool); c=q.flow_candidate.astype(bool)
    tp=int((c&y).sum()); fp=int((c&~y).sum())
    precision=tp/max(tp+fp,1); recall=tp/max(int(y.sum()),1); rate=float(c.mean()) if len(c) else 0
    op=q.opal_candidate.astype(bool)
    op_tp=int((op&y).sum()); op_rec=op_tp/max(int(y.sum()),1); op_prec=op_tp/max(int(op.sum()),1)
    only=int((c&y&~op).sum())
    union=(c|op); union_rec=float((union&y).sum()/max(int(y.sum()),1))
    brier=float(np.mean((q.p_flow_reversal-q.reversal_target)**2)) if len(q) else np.nan
    ll=float(log_loss(q.reversal_target,np.clip(q.p_flow_reversal,1e-6,1-1e-6),labels=[0,1])) if len(q) else np.nan
    return {
        "year":year,"eligible_n":len(q),"true_reversal_n":int(y.sum()),
        "flow_candidate_n":int(c.sum()),"flow_precision":precision,"flow_recall":recall,
        "flow_candidate_rate":rate,"opal_candidate_n":int(op.sum()),"opal_precision":op_prec,
        "opal_recall":op_rec,"flow_only_true_reversal_n":only,"union_recall":union_rec,
        "brier":brier,"logloss":ll
    }

def main():
    raw, listing_count=acquire()
    raw.to_csv(OUT_SRC,index=False)
    cov=coverage(raw)
    cov.to_csv(OUT_COV,index=False)

    feat=source_features(raw)
    panel=pd.read_csv(PANEL)
    a=align(panel,feat)
    if a.empty:
        raise RuntimeError("NO_ALIGNED_PANEL")
    pred=walk(a)
    if pred.empty:
        raise RuntimeError("NO_PREDICTIONS")
    pred=merge_opal(pred)

    dev=pred[pred.year.isin([2023,2024])].copy()
    grid=pd.DataFrame([cand_metrics(dev,t) for t in THRESH_GRID])
    grid.to_csv(OUT_GRID,index=False)
    elig=grid[grid.eligible].copy()

    selected=None; confirm=None; holdout=None
    if elig.empty:
        status="NO_ELIGIBLE_FLOW_PRELIM_OI_THRESHOLD"
    else:
        elig=elig.sort_values(["f2","recall","precision","candidate_rate","threshold"],
                              ascending=[False,False,False,True,False])
        selected=float(elig.iloc[0].threshold)
        s25=period_stats(pred,selected,2025)
        confirm=bool(
            s25["flow_recall"]>s25["opal_recall"]
            and s25["flow_only_true_reversal_n"]>=1
            and s25["flow_precision"]>=.40
            and s25["union_recall"]>s25["opal_recall"]
        )
        status="CONFIRM_PASS" if confirm else "CONFIRM_FAIL"
        if confirm:
            h=period_stats(pred,selected,2026)
            z=pred[(pred.year==2026)&pred.aurora_follows_momentum.astype(bool)].copy()
            z["flow_candidate"]=z.p_flow_reversal>=selected
            miss=(z.reversal_target==1)&(z.v5_pred==z.momentum_up)&(~z.opal_candidate)
            hit=int((miss&z.flow_candidate).sum())
            flow_dir=1-z.momentum_up.astype(int)
            changed=z.flow_candidate & (flow_dir!=z.v5_pred)
            rescue=int((changed&(z.v5_pred!=z.y_up)&(flow_dir==z.y_up)).sum())
            broken=int((changed&(z.v5_pred==z.y_up)&(flow_dir!=z.y_up)).sum())
            h.update({
                "covered_v5_missed_reversal_opal_no_candidate_n":int(miss.sum()),
                "flow_hits_in_covered_v5_missed_opal_no_candidate":hit,
                "diagnostic_v5_rescue":rescue,"diagnostic_v5_broken":broken,
                "diagnostic_v5_net":rescue-broken
            })
            holdout=h

    pred.to_csv(OUT_PRED,index=False)
    summ={
        "schema":"FLOW_PRELIM_OI_H3_V1","status":status,
        "source":{"host":HOST,"dir":DIR,"listing_count":listing_count,
                  "date_min":str(DATE_MIN.date()),"date_max":str(DATE_MAX.date()),
                  "rows":len(raw),"pass_rows":int((raw.parse_status=="PASS").sum())},
        "coverage":cov.to_dict("records"),
        "aligned_rows":len(a),"prediction_rows":len(pred),
        "selected_threshold":selected,"threshold_grid":grid.to_dict("records"),
        "confirmation_2025":period_stats(pred,selected,2025) if selected is not None else None,
        "confirmation_pass":confirm,"holdout_2026":holdout
    }
    OUT_JSON.write_text(json.dumps(summ,indent=2,default=str)+"\n")

    lines=[
        "# FLOW-PRELIM-OI-H3 V1 — RESULT","",
        f"**Status:** **{status}**  ",
        "**Source:** CME anonymous FTP daily_volume XLSX; official CME preliminary GC Volume+OI.  ",
        "**PIT:** strict prior-trade-date only; same-day preliminary data forbidden.","",
        "## Source coverage","",
        "| Year | Listed files | PASS | Coverage | First PASS | Last PASS | Missing |",
        "|---:|---:|---:|---:|---|---|---:|"
    ]
    for r in cov.itertuples():
        lines.append(f"| {r.year} | {r.listed_files} | {r.pass_rows} | {100*r.coverage:.2f}% | {r.min_pass} | {r.max_pass} | {r.missing_rows} |")

    lines += ["","## DEV 2023-2024 threshold grid","",
              "| Threshold | Candidate | Precision | Recall | Rate | F2 | Eligible |",
              "|---:|---:|---:|---:|---:|---:|---|"]
    for r in grid.itertuples():
        lines.append(f"| {r.threshold:.2f} | {r.candidate_n} | {100*r.precision:.2f}% | {100*r.recall:.2f}% | {100*r.candidate_rate:.2f}% | {r.f2:.4f} | {r.eligible} |")

    if selected is not None:
        s25=period_stats(pred,selected,2025)
        lines += ["",f"## Selected threshold: {selected:.2f}","",
                  "## 2025 confirmation","",
                  f"- FLOW recall: **{100*s25['flow_recall']:.2f}%**",
                  f"- OPAL recall same universe: **{100*s25['opal_recall']:.2f}%**",
                  f"- FLOW precision: **{100*s25['flow_precision']:.2f}%**",
                  f"- FLOW-only true OPAL-missed reversals: **{s25['flow_only_true_reversal_n']}**",
                  f"- OPAL ∪ FLOW recall: **{100*s25['union_recall']:.2f}%**",
                  f"- Confirmation: **{'PASS' if confirm else 'FAIL'}**"]
        if confirm and holdout is not None:
            h=holdout
            lines += ["","## 2026 covered holdout","",
                      f"- eligible covered origins: **{h['eligible_n']}**",
                      f"- true reversals: **{h['true_reversal_n']}**",
                      f"- FLOW recall: **{100*h['flow_recall']:.2f}%**",
                      f"- FLOW precision: **{100*h['flow_precision']:.2f}%**",
                      f"- OPAL recall same covered universe: **{100*h['opal_recall']:.2f}%**",
                      f"- OPAL ∪ FLOW recall: **{100*h['union_recall']:.2f}%**",
                      f"- covered V5-missed + OPAL-no-candidate reversals: **{h['covered_v5_missed_reversal_opal_no_candidate_n']}**",
                      f"- FLOW hits inside that covered set: **{h['flow_hits_in_covered_v5_missed_opal_no_candidate']}**",
                      f"- diagnostic forced-route rescue / broken / net: **{h['diagnostic_v5_rescue']} / {h['diagnostic_v5_broken']} / {h['diagnostic_v5_net']:+d}**"]

    lines += ["","## Governance","",
              "This is the separately named preliminary-OI challenger. The original FINAL-only FLOW-H3 V1 remains unchanged. No 2026 outcome selected features, lags, model or threshold."]
    OUT_MD.write_text("\n".join(lines)+"\n")
    print(OUT_MD.read_text())

if __name__=="__main__":
    main()

# trigger: prelim-oi-workflow-ready
